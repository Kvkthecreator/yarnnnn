"""The outbound writer — workspace material written AS an office file.

ADR-395 amendment 2, D16 (phase 3). ADR-417 §2b's demand gate fired: "if
downloadable export returns, it returns as an in-API library call, never as a
standing deployed service". This module is that library call.

── What it is ────────────────────────────────────────────────────────────────

Four deterministic writers, each `(source_text, title) -> bytes`:

  markdown_to_docx   Markdown → a Word document. Markdown is rendered to HTML by
                     the `markdown` library (the same parser `compose/engine.py`
                     uses) and then walked by the ONE structural walker below —
                     so a `.md` and an `.html` become Word the same way.
  html_to_docx       HTML (a Text/Studio artifact) → a Word document.
  deck_to_pptx       A Slides deck's HTML → a PowerPoint deck, one slide per
                     `<section class="slide">` (the deck model: `authoring.py`'s
                     `deck` layout — the first heading block is the slide title,
                     everything else is its body).
  csv_to_xlsx /      A table → one sheet.
  tsv_to_xlsx

WHICH writer serves WHICH source is not decided here. Each writer is bound to
its source format in the format registry (`services/file_formats.py`,
`ExportTarget.writer`) — the registry is the one place a format is named, and a
row carries its writer exactly as it carries its extractor.

── What it deliberately is NOT ───────────────────────────────────────────────

  * A renderer. Styling fidelity is BEST-EFFORT: structure (headings, lists,
    tables, emphasis, links, slide titles) survives; CSS, fonts, colours,
    images and layout do not. The words survive, and the gate proves it by
    reading every output back through the phase-1 extractor.
  * An executor. Nothing in the source runs. HTML is parsed (lxml, network
    off), never evaluated; `<script>`/`<style>` are skipped as non-content;
    a CSV cell that begins with `=` is written as TEXT, never as a formula —
    a spreadsheet must not compute what a member's data merely says.
  * Unbounded. Every writer refuses (`ExportRefused`) past a size cap rather
    than truncating — a silently shortened contract is worse than a refusal.

Module scope imports third-party libraries only (python-docx, python-pptx,
openpyxl, markdown, lxml), because `file_formats` imports this module to bind
the writers and must stay importable from `content_types` without a cycle. The
WRITE — reading a source, naming the sibling, landing the revision, deriving
its projection — is `write_office_file`, whose service imports are local.

Canonical reference:
docs/adr/ADR-395-model-consumable-projection-and-upload-intake-conformance.md §11.11
"""

from __future__ import annotations

import csv
import io
import logging
import re
from typing import Any, Callable, Optional

logger = logging.getLogger(__name__)

#: Largest source text any writer accepts (characters). Bigger than any
#: document a member would hand a client; small enough that one request cannot
#: pin the API building it.
MAX_SOURCE_CHARS = 2_000_000
#: Largest table the sheet writer accepts.
MAX_ROWS = 100_000
MAX_COLS = 1_000
#: Largest deck the slide writer accepts.
MAX_SLIDES = 300
#: Largest file any writer may produce — the upload door's own ceiling, so an
#: export can never mint a file the member could not have uploaded.
MAX_OUTPUT_BYTES = 25 * 1024 * 1024

Writer = Callable[[str, str], bytes]


class ExportRefused(ValueError):
    """The writer declined this source — the message is member-legible."""


def _bounded(text: str) -> str:
    text = text or ""
    if len(text) > MAX_SOURCE_CHARS:
        raise ExportRefused(
            f"The source is {len(text):,} characters — past the {MAX_SOURCE_CHARS:,}-"
            "character limit for writing an office file. Split it into smaller files."
        )
    return text


def _out(buf: io.BytesIO) -> bytes:
    data = buf.getvalue()
    if len(data) > MAX_OUTPUT_BYTES:
        raise ExportRefused("The written file would exceed 25 MB.")
    return data


# ── HTML parsing (shared) ──────────────────────────────────────────────────


#: Elements whose contents are never document text.
_SKIP = frozenset({
    "script", "style", "head", "template", "noscript", "svg", "canvas",
    "iframe", "object", "embed", "button", "input", "select", "textarea",
})
_HEADINGS = {f"h{n}": n for n in range(1, 7)}
_BLOCKS = frozenset({
    "p", "div", "section", "article", "main", "header", "footer", "aside",
    "nav", "figure", "figcaption", "body", "html", "blockquote", "pre",
    "ul", "ol", "li", "table", "hr", "dl", "dt", "dd", "details", "summary",
    *_HEADINGS,
})


def _parse_html(html: str):
    """Parse HTML into an lxml tree — parsing only, network OFF, comments gone."""
    import lxml.html

    parser = lxml.html.HTMLParser(
        remove_comments=True, remove_pis=True, no_network=True, recover=True,
    )
    return lxml.html.document_fromstring(html or "<html></html>", parser=parser)


def _tag(el) -> str:
    t = el.tag
    return t.lower() if isinstance(t, str) else ""


def _text_of(el) -> str:
    """An element's visible text, whitespace collapsed, skipped subtrees out."""
    parts: list[str] = []

    def walk(node):
        if _tag(node) in _SKIP:
            return
        if node.text:
            parts.append(node.text)
        for child in node:
            walk(child)
            if child.tail:
                parts.append(child.tail)

    walk(el)
    return re.sub(r"\s+", " ", "".join(parts)).strip()


def _classes(el) -> set[str]:
    return set((el.get("class") or "").split())


# ── HTML → docx: the ONE structural walker ─────────────────────────────────


def _add_hyperlink(paragraph, url: str, text: str, style: dict) -> None:
    """A real Word hyperlink (python-docx has no public API for one)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    if style.get("bold"):
        props.append(OxmlElement("w:b"))
    if style.get("italic"):
        props.append(OxmlElement("w:i"))
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


_LINKABLE = re.compile(r"^(https?:|mailto:)", re.IGNORECASE)


class _DocxWriter:
    """Walks an HTML tree into a python-docx Document.

    Block elements open paragraphs; inline elements become runs carrying the
    emphasis in force. Loose text directly inside a container is collected
    into a paragraph of its own rather than dropped.
    """

    def __init__(self):
        import docx

        self.doc = docx.Document()
        self.title: Optional[str] = None

    # -- inline -------------------------------------------------------------

    def _runs(self, paragraph, el, style: dict) -> None:
        """Emit el's inline content (text + inline children) into `paragraph`."""
        tag = _tag(el)
        if tag in _SKIP:
            return
        style = dict(style)
        if tag in ("strong", "b"):
            style["bold"] = True
        elif tag in ("em", "i", "cite"):
            style["italic"] = True
        elif tag in ("code", "kbd", "samp", "tt"):
            style["mono"] = True
        elif tag in ("u", "ins"):
            style["underline"] = True
        elif tag in ("s", "del", "strike"):
            style["strike"] = True
        elif tag == "br":
            paragraph.add_run().add_break()
        elif tag == "img":
            alt = (el.get("alt") or "").strip()
            if alt:
                self._run(paragraph, f"[{alt}]", {**style, "italic": True})
            return
        elif tag == "a":
            href = (el.get("href") or "").strip()
            text = _text_of(el)
            if text and _LINKABLE.match(href):
                _add_hyperlink(paragraph, href, text, style)
                return
        if el.text:
            self._run(paragraph, el.text, style)
        for child in el:
            self._runs(paragraph, child, style)
            if child.tail:
                self._run(paragraph, child.tail, style)

    def _run(self, paragraph, text: str, style: dict) -> None:
        text = re.sub(r"\s+", " ", text)
        if not text:
            return
        run = paragraph.add_run(text)
        run.bold = style.get("bold") or None
        run.italic = style.get("italic") or None
        run.underline = style.get("underline") or None
        if style.get("strike"):
            run.font.strike = True
        if style.get("mono"):
            run.font.name = "Courier New"

    @staticmethod
    def _tidy(paragraph) -> None:
        """Trim the whitespace a collapsed run left at either edge."""
        runs = paragraph.runs
        if runs:
            runs[0].text = runs[0].text.lstrip()
            runs[-1].text = runs[-1].text.rstrip()

    # -- blocks -------------------------------------------------------------

    def _paragraph(self, el, style_name: Optional[str] = None) -> None:
        p = self.doc.add_paragraph(style=style_name) if style_name else self.doc.add_paragraph()
        self._runs(p, el, {})
        self._tidy(p)
        if not p.text.strip():
            p._p.getparent().remove(p._p)

    def _list(self, el, depth: int) -> None:
        ordered = _tag(el) == "ol"
        base = "List Number" if ordered else "List Bullet"
        style = base if depth == 0 else f"{base} {min(depth + 1, 3)}"
        for li in el:
            if _tag(li) != "li":
                continue
            p = self.doc.add_paragraph(style=style)
            # The item's own inline content; a nested list becomes deeper items.
            if li.text:
                self._run(p, li.text, {})
            nested = []
            for child in li:
                if _tag(child) in ("ul", "ol"):
                    nested.append(child)
                elif _tag(child) == "p":
                    self._runs(p, child, {})
                else:
                    self._runs(p, child, {})
                if child.tail:
                    self._run(p, child.tail, {})
            self._tidy(p)
            for sub in nested:
                self._list(sub, depth + 1)

    def _table(self, el) -> None:
        rows = [r for r in el.iter("tr")]
        grid = [[c for c in r if _tag(c) in ("td", "th")] for r in rows]
        grid = [r for r in grid if r]
        if not grid:
            return
        cols = max(len(r) for r in grid)
        table = self.doc.add_table(rows=len(grid), cols=cols)
        table.style = "Table Grid"
        for i, row in enumerate(grid):
            for j, cell in enumerate(row):
                target = table.cell(i, j).paragraphs[0]
                self._runs(target, cell, {"bold": _tag(cell) == "th" or None})
                self._tidy(target)

    def block(self, el, depth: int = 0) -> None:
        tag = _tag(el)
        if tag in _SKIP:
            return
        if tag == "title":
            self.title = self.title or _text_of(el) or None
            return
        if tag in _HEADINGS:
            text = _text_of(el)
            if text:
                self.doc.add_heading(text, level=_HEADINGS[tag])
                if tag == "h1" and not self.title:
                    self.title = text
            return
        if tag in ("ul", "ol"):
            self._list(el, depth)
            return
        if tag == "table":
            self._table(el)
            return
        if tag == "pre":
            p = self.doc.add_paragraph()
            for line_no, line in enumerate(el.text_content().splitlines()):
                if line_no:
                    p.add_run().add_break()
                run = p.add_run(line)
                run.font.name = "Courier New"
            return
        if tag == "blockquote":
            self._container(el, depth, quote=True)
            return
        if tag == "hr":
            return
        if tag == "p":
            self._paragraph(el)
            return
        self._container(el, depth)

    def _container(self, el, depth: int, quote: bool = False) -> None:
        """A block container: block children recurse; loose inline content is
        gathered into paragraphs so no text is dropped."""
        pending: list = []  # inline nodes + text awaiting a paragraph

        def flush():
            if not pending:
                return
            p = self.doc.add_paragraph(style="Quote" if quote else None)
            for item in pending:
                if isinstance(item, str):
                    self._run(p, item, {})
                else:
                    self._runs(p, item, {})
                    if item.tail:
                        self._run(p, item.tail, {})
            self._tidy(p)
            if not p.text.strip():
                p._p.getparent().remove(p._p)
            pending.clear()

        if el.text and el.text.strip():
            pending.append(el.text)
        for child in el:
            tag = _tag(child)
            if tag in _BLOCKS or tag in ("table", "title"):
                flush()
                if quote and tag == "p":
                    self._paragraph(child, "Quote")
                else:
                    self.block(child, depth)
                if child.tail and child.tail.strip():
                    pending.append(child.tail)
            elif tag in _SKIP or not tag:
                if child.tail and child.tail.strip():
                    pending.append(child.tail)
            else:
                pending.append(child)
        flush()

    def finish(self, fallback_title: str) -> bytes:
        self.doc.core_properties.title = (self.title or fallback_title or "")[:255]
        buf = io.BytesIO()
        self.doc.save(buf)
        return _out(buf)


def html_to_docx(source: str, title: str = "") -> bytes:
    """HTML → .docx (best-effort styling; every word kept)."""
    tree = _parse_html(_bounded(source))
    writer = _DocxWriter()
    head_title = tree.find(".//title")
    if head_title is not None:
        writer.title = _text_of(head_title) or None
    body = tree.find(".//body")
    writer.block(body if body is not None else tree)
    return writer.finish(title)


def markdown_to_docx(source: str, title: str = "") -> bytes:
    """Markdown → .docx, through HTML, so both sources share one walker."""
    import markdown

    html = markdown.markdown(
        _bounded(source), extensions=["tables", "fenced_code", "sane_lists"],
    )
    return html_to_docx(f"<html><body>{html}</body></html>", title)


# ── deck HTML → pptx ───────────────────────────────────────────────────────


def _slides_of(tree) -> list:
    return [
        el for el in tree.iter("section")
        if "slide" in _classes(el)
    ]


def deck_to_pptx(source: str, title: str = "") -> bytes:
    """A Slides deck → .pptx, one slide per `<section class="slide">`.

    The deck model (`authoring.py`, layout `deck`): the first slide is the
    title slide (kicker + h1 thesis + framing line); every other slide is one
    idea led by an h2, its body in blocks. So: the slide's first h1/h2 is its
    title; on the title slide the rest of the heading area is the subtitle;
    everywhere else the rest is body text, lists kept as bullets and tables as
    tab-separated rows. The deck model carries no speaker notes, so none are
    written — and none are invented.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    tree = _parse_html(_bounded(source))
    slides = _slides_of(tree)
    if not slides:
        raise ExportRefused(
            "This file has no slides (no <section class=\"slide\">), so it cannot "
            "be written as a presentation."
        )
    if len(slides) > MAX_SLIDES:
        raise ExportRefused(f"The deck has {len(slides)} slides — past the {MAX_SLIDES}-slide limit.")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9, the deck's own aspect
    title_layout, content_layout = prs.slide_layouts[0], prs.slide_layouts[1]

    deck_title: Optional[str] = None
    for index, section in enumerate(slides):
        heading = next(
            (h for h in section.iter() if _tag(h) in ("h1", "h2")), None,
        )
        heading_text = _text_of(heading) if heading is not None else ""
        is_title_slide = section.get("data-arrange") == "title" or (
            index == 0 and heading is not None and _tag(heading) == "h1"
        )
        body_lines = _slide_body(section, heading)

        if is_title_slide:
            slide = prs.slides.add_slide(title_layout)
            slide.shapes.title.text = heading_text
            subtitle = slide.placeholders[1]
            subtitle.text_frame.text = "\n".join(text for text, _ in body_lines)
            deck_title = deck_title or heading_text or None
        else:
            slide = prs.slides.add_slide(content_layout)
            if heading_text:
                slide.shapes.title.text = heading_text
            else:  # an untitled slide keeps no empty "Click to add title" box
                t = slide.shapes.title
                t._element.getparent().remove(t._element)
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for n, (text, level) in enumerate(body_lines):
                para = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
                para.text = text
                para.level = min(level, 4)
            if not body_lines:
                ph = slide.placeholders[1]
                ph._element.getparent().remove(ph._element)
            for para in frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(20)

    prs.core_properties.title = (deck_title or title or "")[:255]
    buf = io.BytesIO()
    prs.save(buf)
    return _out(buf)


def _slide_body(section, heading) -> list[tuple[str, int]]:
    """A slide's body as (text, indent level) lines, the title excluded."""
    lines: list[tuple[str, int]] = []

    def walk(el, level: int):
        tag = _tag(el)
        if el is heading or tag in _SKIP:
            return
        if tag in ("ul", "ol"):
            for li in el:
                if _tag(li) != "li":
                    continue
                own = _text_of_without_lists(li)
                if own:
                    lines.append((own, level))
                for sub in li:
                    if _tag(sub) in ("ul", "ol"):
                        walk(sub, level + 1)
            return
        if tag == "table":
            for tr in el.iter("tr"):
                cells = [_text_of(c) for c in tr if _tag(c) in ("td", "th")]
                if any(cells):
                    lines.append(("\t".join(cells), level))
            return
        if tag in ("p", "figcaption", "blockquote", "pre", *_HEADINGS) or (
            tag in ("div", "span") and not any(_tag(c) in _BLOCKS for c in el)
        ):
            text = _text_of(el)
            if text:
                lines.append((text, level))
            return
        for child in el:
            walk(child, level)

    for child in section:
        walk(child, 0)
    return lines


def _text_of_without_lists(li) -> str:
    parts = [li.text or ""]
    for child in li:
        if _tag(child) not in ("ul", "ol"):
            parts.append(_text_of(child))
        parts.append(child.tail or "")
    return re.sub(r"\s+", " ", " ".join(parts)).strip()


# ── table text → xlsx ──────────────────────────────────────────────────────


_INT = re.compile(r"^-?(0|[1-9]\d{0,14})$")
_FLOAT = re.compile(r"^-?(0|[1-9]\d{0,14})\.\d+$")


def _cell_value(raw: str):
    """A number when the text IS a plain number; otherwise the text, verbatim.

    Deliberately narrow: `007` (an identifier) and `1,200` (a formatted figure)
    stay text, because reading them as numbers would change what they say.
    """
    s = raw.strip()
    if _INT.match(s):
        return int(s)
    if _FLOAT.match(s):
        return float(s)
    return raw


def _table_to_xlsx(source: str, title: str, delimiter: str) -> bytes:
    import openpyxl

    rows = list(csv.reader(io.StringIO(_bounded(source)), delimiter=delimiter))
    if len(rows) > MAX_ROWS:
        raise ExportRefused(f"The table has {len(rows):,} rows — past the {MAX_ROWS:,}-row limit.")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[\\/*?:\[\]]", " ", title).strip() or "Sheet1")[:31]
    for r, row in enumerate(rows, start=1):
        if len(row) > MAX_COLS:
            raise ExportRefused(f"Row {r} has {len(row):,} columns — past the {MAX_COLS:,}-column limit.")
        for c, raw in enumerate(row, start=1):
            if raw == "":
                continue
            cell = ws.cell(row=r, column=c, value=_cell_value(raw))
            # NEVER a formula. openpyxl reads a leading `=` as one; a member's
            # data saying `=HYPERLINK(...)` must stay what it says.
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return _out(buf)


def csv_to_xlsx(source: str, title: str = "") -> bytes:
    """CSV → .xlsx, one sheet."""
    return _table_to_xlsx(source, title, ",")


def tsv_to_xlsx(source: str, title: str = "") -> bytes:
    """TSV → .xlsx, one sheet."""
    return _table_to_xlsx(source, title, "\t")


# ── THE WRITE — a new attributed file, derived from its source ────────────


def _leaf_stem(path: str) -> str:
    leaf = path.rsplit("/", 1)[-1]
    return leaf.rsplit(".", 1)[0] if "." in leaf else leaf


def sibling_export_path(source_abs: str, to: str, taken: set[str]) -> Optional[str]:
    """Where a member's "Save as" lands: beside its source, never over a file.

    `{dir}/{stem}.{to}`, then `{stem}-2.{to}`, `-3`… — the same `-N` grammar
    the upload door uses for a name already taken (`documents._unique_raw_path`).
    An export is always a NEW file; it never replaces one (ADR-395 §11.11).
    """
    head = source_abs.rpartition("/")[0]
    stem = _leaf_stem(source_abs)
    for n in range(1, 201):
        leaf = f"{stem}.{to}" if n == 1 else f"{stem}-{n}.{to}"
        candidate = f"{head}/{leaf}"
        if candidate not in taken:
            return candidate
    return None


def _looks_like_html(text: str) -> bool:
    return bool(re.match(r"\s*<", text or ""))


async def write_office_file(
    auth: Any,
    *,
    target_path: str,
    source_text: Optional[str],
    derived_from: Optional[list],
    authored_by: str,
    message: str,
    author_identity_uuid: Optional[str] = None,
    expected_parent_version_id: Any = None,
) -> dict:
    """Write `target_path` (an office path) as a NEW attributed revision.

    The ONE write both doors reach: the member's "Save as" (the route) and an
    agent's WriteFile to an office path (the lane, and MCP `save`, which
    dispatches WriteFile). Two source forms:

      * `source_text` given — the content IS the source: Markdown or HTML for a
        document, a Slides deck's HTML for a presentation, CSV for a sheet.
      * `source_text` None and exactly one `derived_from` — the kernel reads
        that file's head and converts it. No bytes cross the model: a 40-slide
        deck is not re-emitted token by token to be exported.

    Either way the format registry decides whether that source may become this
    target (`file_formats.export_writer`), the revision lands through
    `write_revision` (binary lane, ADR-427) with `derived_from` recorded, and
    the output's text projection is derived exactly as an upload's is — so the
    file an export mints is one yarnnn reads back.
    """
    from services.authored_substrate import normalize_workspace_ref, write_revision
    from services.file_formats import ext_of, export_writer, inline_source_ext

    to = ext_of(target_path)
    sources = [s for s in (normalize_workspace_ref(p) for p in (derived_from or [])) if s]

    source_path: Optional[str] = None
    if source_text is None:
        if len(sources) != 1:
            return {
                "success": False, "error": "export_source_required",
                "message": (
                    f"To write a .{to} from an existing file, name exactly one source in "
                    "derived_from — or pass the source text as content."
                ),
            }
        source_path = sources[0]
        loaded = _read_source(auth, source_path)
        if "error" in loaded:
            return {"success": False, **loaded}
        source_text = loaded["content"]
        src_ext = ext_of(source_path)
        probe_path = source_path
    else:
        src_ext = inline_source_ext(to, _looks_like_html(source_text))
        probe_path = f"inline.{src_ext}" if src_ext else ""

    writer = export_writer(probe_path, to, source_text) if probe_path else None
    if writer is None:
        return {"success": False, "error": "export_not_offered", "message": _not_offered(to, source_path)}

    title = _leaf_stem(source_path or target_path)
    try:
        data = writer(source_text, title)
    except ExportRefused as refused:
        return {"success": False, "error": "export_refused", "message": str(refused)}
    except Exception as exc:  # noqa: BLE001 — a library fault is a refusal, never a 500
        logger.exception("[EXPORT] writer failed for %s", target_path)
        return {"success": False, "error": "export_failed", "message": f"Could not write the .{to} file: {exc}"}

    kwargs: dict = {}
    if expected_parent_version_id is not None:
        kwargs["expected_parent_version_id"] = expected_parent_version_id
    revision_id = write_revision(
        auth.client,
        user_id=auth.user_id,
        path=target_path,
        content_bytes=data,
        authored_by=authored_by,
        author_identity_uuid=author_identity_uuid,
        message=message,
        summary=(f"Written as .{to} from {source_path}" if source_path else f"Written as .{to}"),
        lifecycle="active",
        workspace_id=getattr(auth, "workspace_id", None),
        revision_kind="derivation" if sources else "authored",
        derived_from=sources or None,
        **kwargs,
    )

    projection = await _derive_projection(auth, target_path, data, to)
    return {
        "success": True,
        "path": target_path,
        "revision_id": revision_id,
        "format": to,
        "bytes": len(data),
        "derived_from": sources,
        **projection,
    }


def _not_offered(to: str, source_path: Optional[str]) -> str:
    from services.file_formats import FORMATS

    makers = sorted({
        f".{fmt.exts[0]}" + (" (a Slides deck)" if t.when_app == "slides" else "")
        for fmt in FORMATS for t in fmt.export_as if t.to == to
    })
    what = f"`{source_path}`" if source_path else "this content"
    if not makers:
        return f"yarnnn does not write .{to} files."
    return f"{what} cannot be written as .{to}. A .{to} is made from: {', '.join(makers)}."


def _read_source(auth: Any, source_path: str) -> dict:
    """The source's head TEXT, under the caller's read grant."""
    from services.primitives.workspace import _is_path_readable_for_principal, _scope_filter

    if not _is_path_readable_for_principal(auth, source_path):
        return {"error": "source_not_readable",
                "message": f"Your grant does not permit reading {source_path}."}
    rows = (
        auth.client.table("workspace_files")
        .select("content, content_type")
        .eq(*_scope_filter(auth))
        .eq("path", source_path)
        .limit(1)
        .execute()
    ).data or []
    if not rows:
        return {"error": "source_not_found", "message": f"No live file at {source_path}."}
    content = rows[0].get("content") or ""
    if not content.strip():
        return {"error": "source_empty",
                "message": f"{source_path} has no text to write out."}
    return {"content": content}


async def _derive_projection(auth: Any, raw_path: str, data: bytes, file_type: str) -> dict:
    """The output's text projection — the SAME derive an upload runs.

    Reads the written bytes back through the phase-1 extractor (not the source
    text: what is projected is what the file actually says) and lands the
    co-located `.extracted.md` via ExtractTextFromBlob, attributed
    `system:extract`. Best-effort: the office file is already landed and is
    the file that matters; a missing projection is a missing index entry.
    """
    from services.documents import derive_upload_projection, extract_text

    try:
        text, _units = await extract_text(data, file_type)
        return await derive_upload_projection(
            auth.client, auth.user_id, raw_path, text,
            filename=raw_path.rsplit("/", 1)[-1], file_type=file_type,
            workspace_id=getattr(auth, "workspace_id", None),
        )
    except Exception as exc:  # noqa: BLE001
        logger.warning("[EXPORT] projection derive failed for %s: %s", raw_path, exc)
        return {"projection_path": None, "word_count": 0}


__all__ = [
    "ExportRefused",
    "MAX_OUTPUT_BYTES",
    "MAX_SOURCE_CHARS",
    "csv_to_xlsx",
    "deck_to_pptx",
    "html_to_docx",
    "markdown_to_docx",
    "sibling_export_path",
    "tsv_to_xlsx",
    "write_office_file",
]
