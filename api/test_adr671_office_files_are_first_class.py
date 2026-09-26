"""
ADR-671 — Office files are first-class: the member's format is the document.

Drives the office kernel on REAL files (written by python-docx / python-pptx /
openpyxl) and asserts on the package it writes back, never on a mock of it:

  D2  every address the projection prints resolves to the element it names;
  D3  an edit lands exactly the change asked, and every other part of the
      package has identical content (the loss this ADR exists to stop);
  D4  a Word edit authored through a model is a tracked change signed by the
      principal's display name; a member's own edit is clean;
  D10 formulas survive, a shared-formula master keeps its dependents, DDE is
      refused, a merged cell names its anchor, a formula edit drops calcChain;
  D1  an existing office file is never rebuilt from a source;
  and the doors: EditFile, MCP `edit` and `open`, ReadFile's guidance, the
  restore route and the client slide view (executed across the language line).
"""

import asyncio
import io
import json
import re
import shutil
import subprocess
import types
import zipfile
from pathlib import Path
from unittest.mock import patch

import pytest

ROOT = Path(__file__).resolve().parent

docx_lib = pytest.importorskip("docx")
pptx_lib = pytest.importorskip("pptx")
openpyxl = pytest.importorskip("openpyxl")

from services.office import docx as odocx  # noqa: E402
from services.office import pptx as opptx  # noqa: E402
from services.office import xlsx as oxlsx  # noqa: E402
from services.office.package import OfficeEditError  # noqa: E402


# ── fixtures: real files ────────────────────────────────────────────────────


def _docx_bytes() -> bytes:
    d = docx_lib.Document()
    d.add_heading("Quarterly Report", 1)
    p = d.add_paragraph("Revenue grew to ")
    p.add_run("1,100").bold = True
    p.add_run(" units this quarter.")
    t = d.add_table(rows=2, cols=2)
    for (r, c), v in {(0, 0): "Item", (0, 1): "Q3", (1, 0): "Rent", (1, 1): "1,100"}.items():
        t.cell(r, c).text = v
    d.add_paragraph("Payment is due in 30 days.")
    d.sections[0].header.paragraphs[0].text = "ACME Confidential"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Q3", "Share"])
    for i, (n, v) in enumerate([("Rent", 1100), ("Payroll", 2000), ("Tools", 300)], start=2):
        ws.append([n, v, None])
        ws[f"C{i}"] = f"=B{i}/B5"
        ws[f"C{i}"].number_format = "0%"
    ws["A5"], ws["B5"] = "Total", "=SUM(B2:B4)"
    ws.merge_cells("A7:C7")
    ws["A7"] = "Notes"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    ws.add_chart(chart, "E2")
    wb.create_sheet("Q3 Plan")["A1"] = "note"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx.util import Inches, Pt

    prs = pptx_lib.Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Q3 Plan"
    tf = s1.placeholders[1].text_frame
    tf.text = "Grow revenue by 12%"
    tf.add_paragraph().text = "Hire two engineers"
    tf.paragraphs[0].runs[0].font.size = Pt(28)
    s1.notes_slide.notes_text_frame.text = "Emphasize hiring"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Numbers"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for (r, c), v in {(0, 0): "Metric", (0, 1): "Value", (1, 0): "ARR", (1, 1): "$2.1M"}.items():
        tbl.cell(r, c).text = v
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _parts(data: bytes) -> dict:
    z = zipfile.ZipFile(io.BytesIO(data))
    return {n: z.read(n) for n in z.namelist()}


def _changed(before: bytes, after: bytes) -> set:
    a, b = _parts(before), _parts(after)
    return {n for n in a if n not in b or a[n] != b[n]} | (set(b) - set(a))


def _xml(data: bytes, part: str) -> str:
    return _parts(data)[part].decode("utf-8")


# ── D2 — the address a reader sees is the address an edit resolves ──────────


def test_d2_every_docx_address_resolves_to_the_paragraph_it_labels():
    data = _docx_bytes()
    text, _ = odocx.project(data)
    labels = re.findall(r"\[p(\d+)(?: · [^\]]+)?\] ([^|\n]+?)(?= \||\n|$)", text)
    assert len(labels) >= 7, text
    body = docx_lib.Document(io.BytesIO(data)).element.body
    paras = odocx._paragraphs(body)
    for n, shown in labels:
        assert odocx._text(paras[int(n) - 1]).strip() == shown.strip(), (n, shown)
    assert "[p1 · Heading 1] Quarterly Report" in text
    assert "Header: ACME Confidential" in text


def test_d2_every_xlsx_cell_is_addressed_and_shows_formula_beside_value():
    data = _xlsx_bytes()
    # openpyxl writes no cached values; give B5 one, as Excel would.
    parts = _parts(data)
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    sheet = re.sub(r'(<c r="B5"[^>]*>)<f>SUM\(B2:B4\)</f><v\s*/>|(<c r="B5"[^>]*>)<f>SUM\(B2:B4\)</f>',
                   lambda m: (m.group(1) or m.group(2)) + "<f>SUM(B2:B4)</f><v>3400</v>", sheet)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, b in parts.items():
            z.writestr(n, sheet.encode() if n == "xl/worksheets/sheet1.xml" else b)
    text, sheets = oxlsx.project(buf.getvalue())
    assert sheets == 2
    assert "|  | A | B | C |" in text, "the column-letter header is the address's other half"
    assert "| 5 | Total | =SUM(B2:B4) → 3400 |" in text
    assert "| 2 | Rent | 1100 | =B2/B5 |" in text
    assert "## Q3 Plan" in text


def test_d2_every_pptx_address_resolves_to_the_shape_it_labels():
    text, _ = opptx.project(_pptx_bytes())
    assert "## Slide 1" in text and "[s1/notes] Speaker notes: Emphasize hiring" in text
    assert re.search(r"\[s1/\d+ · Title\] Q3 Plan", text)
    body = re.search(r"\[(s1/\d+)\] Grow revenue", text).group(1)
    cell = re.search(r"\[(s2/\d+/r2c2)\] \$2\.1M", text).group(1)
    out, _ = opptx.apply(_pptx_bytes(), [{"op": "replace", "at": body, "old": "12%", "new": "15%"},
                                         {"op": "replace", "at": cell, "new": "$2.4M"}])
    after, _ = opptx.project(out)
    assert "Grow revenue by 15%" in after and "$2.4M" in after


# ── D3 — exactly the change asked, and nothing else ─────────────────────────


def test_d3_docx_edit_touches_only_the_document_part_and_keeps_run_formatting():
    data = _docx_bytes()
    out, done = odocx.apply(data, [{"op": "replace", "at": "p2", "old": "1,100", "new": "1,250"}])
    assert done == ["replaced text in p2"]
    assert _changed(data, out) == {"word/document.xml"}
    xml = _xml(out, "word/document.xml")
    i = xml.index("1,250")
    assert "<w:b/>" in xml[max(0, i - 300):i], "the new words lost the bold of the run they replaced"
    assert "Revenue grew to 1,250 units this quarter." in odocx.project(out)[0]


def test_d3_xlsx_edit_keeps_chart_styles_formulas_and_the_cell_style():
    data = _xlsx_bytes()
    out, _ = oxlsx.apply(data, [{"op": "replace", "at": "Budget!B2", "old": "1100", "new": "1,250"}])
    changed = _changed(data, out)
    assert changed <= {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}, changed
    assert "xl/charts/chart1.xml" in _parts(out), "the chart must survive an edit"
    wb = openpyxl.load_workbook(io.BytesIO(out))
    ws = wb["Budget"]
    assert ws["B2"].value == 1250 and ws["B5"].value == "=SUM(B2:B4)"
    assert ws["C2"].value == "=B2/B5" and ws["C2"].number_format == "0%"
    assert 'fullCalcOnLoad="1"' in _xml(out, "xl/workbook.xml"), "Excel must recalculate on open"
    assert oxlsx.RECALC_NOTE in oxlsx.project(out)[0]


def test_d3_a_cell_written_past_the_sheets_extent_is_read_back():
    """openpyxl's reader trusts `<dimension>`; a write outside it must widen it,
    or the projection silently omits the member's new value."""
    out, _ = oxlsx.apply(_xlsx_bytes(), [{"op": "replace", "at": "'Q3 Plan'!C3", "new": "hello"}])
    assert "| 3 |  |  | hello |" in oxlsx.project(out)[0]
    assert re.search(r'<dimension ref="A1:C3"', _xml(out, "xl/worksheets/sheet2.xml"))


def test_d2_a_stale_stored_dimension_from_another_tool_hides_nothing():
    """Another tool can leave `<dimension ref="A1">` over a sheet that reaches C3;
    the projection reads the rows, not the hint."""
    parts = _parts(_xlsx_bytes())
    sheet = parts["xl/worksheets/sheet2.xml"].decode()
    sheet = sheet.replace("</sheetData>", '<row r="3"><c r="C3" t="inlineStr"><is><t>late</t></is></c></row></sheetData>')
    assert re.search(r'<dimension ref="A1(:A1)?"\s*/>', sheet), "the fixture must understate its extent"
    parts["xl/worksheets/sheet2.xml"] = sheet.encode()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    assert "| 3 |  |  | late |" in oxlsx.project(buf.getvalue())[0]


def test_d3_pptx_edit_keeps_the_run_font_and_touches_only_that_slide():
    data = _pptx_bytes()
    text, _ = opptx.project(data)
    body = re.search(r"\[(s1/\d+)\] Grow revenue", text).group(1)
    out, _ = opptx.apply(data, [{"op": "replace", "at": body, "old": "12%", "new": "15%"}])
    assert _changed(data, out) == {"ppt/slides/slide1.xml"}
    run = pptx_lib.Presentation(io.BytesIO(out)).slides[0].placeholders[1].text_frame.paragraphs[0].runs[0]
    assert run.text == "Grow revenue by 15%" and run.font.size.pt == 28


# ── D4 — attribution travels inside the Word file ───────────────────────────


def test_d4_a_model_authored_word_edit_is_a_tracked_change_signed_by_the_principal():
    data = _docx_bytes()
    who = "Kevin via Claude Sonnet 5"
    out, _ = odocx.apply(data, [
        {"op": "replace", "at": "p2", "old": "1,100", "new": "1,250"},
        {"op": "insert", "at": "p2", "new": "A new clause.", "style": "Heading 2"},
        {"op": "replace", "at": "p7", "new": ""},
        {"op": "style", "at": "p1", "style": "Title"},
    ], author=who)
    xml = _xml(out, "word/document.xml")
    for tag in ("ins", "del", "pPrChange"):
        assert re.search(rf'<w:{tag} [^>]*w:author="{who}"', xml), f"no tracked w:{tag} by {who}"
    ids = re.findall(r'<w:(?:ins|del|pPrChange) w:id="(\d+)"', xml)
    assert len(ids) == len(set(ids)), "annotation ids must be unique"
    assert "<w:delText>1,100</w:delText>" in xml
    text, _ = odocx.project(out)
    assert text.startswith(odocx.TRACKED_NOTE)
    assert "1,250" in text and "Payment is due" not in text, "the projection reads as if accepted"
    docx_lib.Document(io.BytesIO(out))  # still a document python-docx opens


def test_d4_a_members_own_word_edit_is_clean():
    out, _ = odocx.apply(_docx_bytes(), [{"op": "replace", "at": "p2", "old": "1,100", "new": "1,250"}])
    xml = _xml(out, "word/document.xml")
    assert "<w:ins " not in xml and "<w:del " not in xml


def test_d4_paragraph_mark_revision_is_the_first_child_of_its_rpr():
    """CT_ParaRPr puts ins/del FIRST — Word treats an out-of-order child as damage.
    The paragraph's mark is already formatted (bold), so an APPENDED marker would
    land after `w:b` — the case that makes this arm able to fail."""
    from lxml import etree

    d = docx_lib.Document(io.BytesIO(_docx_bytes()))
    mark = d.paragraphs[-1]._p.get_or_add_pPr()
    rpr = etree.SubElement(mark, odocx.w("rPr"))
    etree.SubElement(rpr, odocx.w("b"))
    buf = io.BytesIO()
    d.save(buf)
    out, _ = odocx.apply(buf.getvalue(), [{"op": "replace", "at": "p7", "new": ""}], author="A")
    root = etree.fromstring(_parts(out)["word/document.xml"])
    marks = [r for r in root.iter(odocx.w("rPr")) if r.getparent().tag == odocx.w("pPr")
             and r.find(odocx.w("del")) is not None]
    assert marks, "the deleted paragraph's mark was not marked deleted"
    assert all(r[0].tag == odocx.w("del") for r in marks), "the deletion must lead its w:rPr"


# ── a stale address fails loudly, never lands on a neighbour ────────────────


@pytest.mark.parametrize("mod,data_fn,edit,needle", [
    (odocx, _docx_bytes, {"op": "replace", "at": "p2", "old": "9,999", "new": "x"}, "does not contain"),
    (odocx, _docx_bytes, {"op": "replace", "at": "p99", "new": "x"}, "does not exist"),
    (oxlsx, _xlsx_bytes, {"op": "replace", "at": "Budget!B2", "old": "999", "new": "1"}, "holds"),
    (oxlsx, _xlsx_bytes, {"op": "replace", "at": "B2", "new": "1"}, "2 sheets"),
    (oxlsx, _xlsx_bytes, {"op": "replace", "at": "Nope!A1", "new": "1"}, "no sheet"),
    (opptx, _pptx_bytes, {"op": "replace", "at": "s9/2", "new": "x"}, "does not exist"),
    (opptx, _pptx_bytes, {"op": "replace", "at": "s1/999", "new": "x"}, "no shape"),
])
def test_a_stale_or_wrong_address_is_refused(mod, data_fn, edit, needle):
    with pytest.raises(OfficeEditError) as e:
        mod.apply(data_fn(), [edit])
    assert needle in str(e.value)


# ── D10 — the edge cases that break workbooks ───────────────────────────────


def test_d10_dde_and_xlm_formulas_are_refused():
    for bad in ("=cmd|' /C calc'!A0", "=CALL(\"kernel32\",\"WinExec\")", "=EXEC(\"x\")"):
        with pytest.raises(OfficeEditError):
            oxlsx.apply(_xlsx_bytes(), [{"op": "replace", "at": "Budget!B2", "new": bad}])
    oxlsx.apply(_xlsx_bytes(), [{"op": "replace", "at": "Budget!D2", "new": '=IF(B2>1,"a|b","c")'}])


def test_d10_a_merged_cell_names_its_anchor():
    with pytest.raises(OfficeEditError) as e:
        oxlsx.apply(_xlsx_bytes(), [{"op": "replace", "at": "Budget!B7", "new": "x"}])
    assert "A7" in str(e.value)


def _with_shared_formula_and_calc_chain(data: bytes) -> bytes:
    """Rewrite C2:C4 as ONE shared formula (Excel's fill-down) and add a calcChain."""
    parts = _parts(data)
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    sheet = re.sub(r'(<c r="C2"[^>]*>)<f>B2/B5</f>', r'\1<f t="shared" ref="C2:C4" si="0">B2/$B$5</f>', sheet)
    sheet = re.sub(r'(<c r="C3"[^>]*>)<f>B3/B5</f>', r'\1<f t="shared" si="0"/>', sheet)
    sheet = re.sub(r'(<c r="C4"[^>]*>)<f>B4/B5</f>', r'\1<f t="shared" si="0"/>', sheet)
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    parts["xl/calcChain.xml"] = (b'<?xml version="1.0" encoding="UTF-8"?><calcChain xmlns='
                                 b'"http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                                 b'<c r="C2" i="1"/></calcChain>')
    rels = parts["xl/_rels/workbook.xml.rels"].decode().replace(
        "</Relationships>",
        '<Relationship Id="rIdCC" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships/calcChain" Target="calcChain.xml"/></Relationships>')
    parts["xl/_rels/workbook.xml.rels"] = rels.encode()
    ct = parts["[Content_Types].xml"].decode().replace(
        "</Types>",
        '<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-'
        'officedocument.spreadsheetml.calcChain+xml"/></Types>')
    parts["[Content_Types].xml"] = ct.encode()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    return buf.getvalue()


def test_d10_overwriting_a_shared_formula_master_keeps_every_dependent():
    data = _with_shared_formula_and_calc_chain(_xlsx_bytes())
    out, _ = oxlsx.apply(data, [{"op": "replace", "at": "Budget!C2", "new": "0.5"}])
    ws = openpyxl.load_workbook(io.BytesIO(out))["Budget"]
    assert ws["C2"].value == 0.5
    assert ws["C3"].value == "=B3/$B$5" and ws["C4"].value == "=B4/$B$5", (ws["C3"].value, ws["C4"].value)
    assert "xl/calcChain.xml" not in _parts(out), "a formula change must drop the stale calcChain"
    assert "calcChain" not in _xml(out, "xl/_rels/workbook.xml.rels")
    assert "calcChain" not in _xml(out, "[Content_Types].xml")


def test_d10_values_parse_as_a_spreadsheet_reads_them():
    assert oxlsx.parse_value("1,250") == ("number", 1250)
    assert oxlsx.parse_value("25%") == ("number", 0.25)
    assert oxlsx.parse_value("007")[0] == "number"  # typed into a cell, Excel reads 7 too
    assert oxlsx.parse_value("=SUM(A1:A3)") == ("formula", "SUM(A1:A3)")
    assert oxlsx.parse_value("Rent") == ("text", "Rent")
    assert oxlsx.parse_value("")[0] == "clear"
    # Excel's literal-text apostrophe, both ways: a text cell that SAYS a
    # formula never reads as one, and writing it back keeps it text.
    assert oxlsx.show("=HYPERLINK(1)") == "'=HYPERLINK(1)"
    assert oxlsx.parse_value("'=HYPERLINK(1)") == ("text", "=HYPERLINK(1)")


# ── D1 — an existing office file is never rebuilt from a source ─────────────


class _Rows:
    def __init__(self, rows):
        self.rows = rows

    def __getattr__(self, _name):
        return lambda *a, **k: self

    def execute(self):
        return types.SimpleNamespace(data=self.rows)


def test_d1_create_refuses_a_path_that_holds_an_office_file():
    from services.office.create import create_office_file

    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    auth = types.SimpleNamespace(client=_Rows([{"content_type": mime}]), user_id="u1", workspace_id="w1")
    out = asyncio.run(create_office_file(
        auth, target_path="/workspace/deals/contract.docx", source_text="# New", derived_from=None,
        authored_by="operator", message="m",
    ))
    assert out["success"] is False and out["error"] == "office_file_exists"
    assert "EditFile" in out["message"]


def test_d1_text_under_an_office_name_is_not_refused():
    """Markdown stored as `memo.docx` before am.2 is not an office file — writing
    the real format over it loses nothing."""
    from services.office.create import _existing_office_head

    auth = types.SimpleNamespace(client=_Rows([{"content_type": "text/markdown"}]), user_id="u1")
    assert _existing_office_head(auth, "/workspace/memo.docx") is None


# ── the doors ───────────────────────────────────────────────────────────────


def _drive_edit_file(authored_by: str, request: dict, data: bytes, path="deals/contract.docx"):
    """The REAL handle_edit_file → edit_office_file → engine, with the substrate,
    storage and landing edges stubbed. Returns (result, landed_kwargs)."""
    import services.authored_substrate as sub
    import services.documents as docs
    import services.primitives.workspace as pw
    import services.principal_display as pd
    import services.storage_backend as sb
    import services.supabase as supa

    landed: dict = {}
    head = types.SimpleNamespace(id="v-head", blob_sha="sha-1", content="")

    class _Backend:
        def get_blob(self, sha, *, workspace_id=None):
            return data

    async def _land(auth, **kw):
        landed.update(kw)
        return {"revision_id": "v-new", "projection_path": None, "word_count": 0, "embed_pending": False}

    auth = types.SimpleNamespace(client=object(), user_id="u1", workspace_id="w1",
                                 caller_identity=authored_by)
    with patch.object(sub, "read_revision", lambda *a, **k: head), \
         patch.object(sub, "read_head_revision_id", lambda *a, **k: "v-head"), \
         patch.object(sb, "get_storage_backend", lambda _c: _Backend()), \
         patch.object(supa, "get_service_client", lambda: object()), \
         patch.object(docs, "land_binary_revision", _land), \
         patch.object(pd, "resolve_member_names", lambda _c, ids: {i: "Kevin" for i in ids}):
        out = asyncio.run(pw.handle_edit_file(auth, {"path": path, **request}))
    return out, landed


def test_door_editfile_through_a_model_lands_tracked_changes_by_its_display_name():
    out, landed = _drive_edit_file(
        "member:u1 via anthropic/claude-sonnet-5",
        {"anchor": {"at": "p2"}, "old_string": "1,100", "new_string": "1,250"},
        _docx_bytes(),
    )
    assert out["success"] is True and out["tracked_changes"] is True, out
    assert landed["expected_parent_version_id"] == "v-head", "the edit must be conditional on the head it read"
    xml = _xml(landed["data"], "word/document.xml")
    assert re.search(r'<w:ins [^>]*w:author="Kevin via [^"]+"', xml), "tracked under the principal's display name"


def test_door_editfile_by_the_member_is_clean_and_batches_are_one_revision():
    out, landed = _drive_edit_file(
        "operator",
        {"edits": [{"at": "p2", "old_string": "1,100", "new_string": "1,250"},
                   {"after": "p7", "new_string": "Late fees apply."}]},
        _docx_bytes(),
    )
    assert out["success"] is True and out["tracked_changes"] is False
    assert out["replacements"] == 2
    xml = _xml(landed["data"], "word/document.xml")
    assert "<w:ins " not in xml and "Late fees apply." in xml


def test_door_editfile_refusal_is_legible_and_lands_nothing():
    out, landed = _drive_edit_file(
        "operator", {"anchor": {"at": "Budget!B7"}, "new_string": "x"}, _xlsx_bytes(), path="q3.xlsx",
    )
    assert out["success"] is False and out["error"] == "office_edit_refused"
    assert "merged" in out["message"] and not landed


def test_door_editfile_on_an_office_file_without_an_address_is_refused_with_the_grammar():
    out, landed = _drive_edit_file("operator", {"old_string": "1,100", "new_string": "1,250"}, _docx_bytes())
    assert out["success"] is False and "anchor={'at'" in out["message"] and not landed


def test_door_mcp_edit_carries_at_to_editfile():
    import services.mcp_composition as mc
    import services.primitives.registry as reg

    seen: dict = {}

    async def _exec(auth, name, request):
        seen.update(name=name, request=request)
        return {"success": True, "path": "/workspace/q3.xlsx", "message": "1 change(s)", "replacements": 1}

    with patch.object(reg, "execute_primitive", _exec):
        out = asyncio.run(mc.compose_edit(object(), "q3.xlsx", old="", new="1250", at="Budget!B2"))
    assert seen["name"] == "EditFile" and seen["request"]["anchor"] == {"at": "Budget!B2"}
    assert out["success"] is True


def test_door_mcp_open_serves_an_office_files_addressed_words():
    import services.mcp_composition as mc
    import services.primitives.registry as reg
    import services.primitives.workspace as pw

    words = "[p1 · Heading 1] Quarterly Report\n\n[p2] Revenue grew."
    auth = types.SimpleNamespace(client=_Rows([{"path": "/workspace/c.docx", "content": "",
                                                 "updated_at": "t", "content_type": "x"}]),
                                 user_id="u1", workspace_id="w1")

    async def _exec(*_a, **_k):
        return {"revisions": []}

    with patch.object(mc, "resolve_binary_head",
                      lambda *_a: {"blob_sha": "s", "byte_size": 10, "content_type": "application/docx"}), \
         patch.object(mc, "mint_binary_url", lambda *_a: None), \
         patch.object(mc, "live_files_filter", lambda q: q), \
         patch.object(mc, "_substrate_scope", lambda _a: ("user_id", "u1")), \
         patch.object(pw, "_projection_words", lambda *_a: words), \
         patch.object(reg, "execute_primitive", _exec):
        out = asyncio.run(mc.compose_open(auth, "c.docx"))
    assert out["binary"] is True and out["content"] == words, out
    assert out["complete_for_write"] is False
    assert "`at`" in out["explanation"] and "edit" in out["explanation"]


def test_door_readfile_names_the_in_place_loop_and_never_the_rebuild():
    from services.primitives.workspace import _readable_binary_answer

    for path in ("/workspace/a.docx", "/workspace/a.xlsx", "/workspace/a.pptx"):
        msg = _readable_binary_answer("workspace", path, path, "x", 10, "[p1] words", 0)["message"]
        assert "EditFile" in msg and "anchor={'at'" in msg, msg
        assert "WriteFile" not in msg, "the rebuild must never be offered for an office file"
    pdf = _readable_binary_answer("workspace", "a.pdf", "/workspace/a.pdf", "x", 1, "w", 0)["message"]
    assert "does not edit" in pdf


def test_restore_lands_through_the_one_helper():
    src = (ROOT / "routes" / "workspace.py").read_text()
    body = src[src.index("async def restore_binary_revision("):src.index('@router.get("/workspace/revisions"')]
    assert "land_binary_revision(" in body
    assert "write_revision(" not in body and "derive_upload_projection(" not in body


def test_create_lands_through_the_one_helper():
    src = (ROOT / "services" / "office" / "create.py").read_text()
    body = src[src.index("async def create_office_file("):src.index("def _existing_office_head(")]
    assert "land_binary_revision(" in body and "write_revision(" not in body


# ── one home, and what was deleted stays deleted ────────────────────────────


def test_registry_projector_and_editor_come_from_one_module_per_format():
    from importlib import import_module
    from services.file_formats import format_for

    for ext in ("docx", "xlsx", "pptx"):
        fmt = format_for(ext)
        kind = import_module(f"services.office.{ext}").KIND
        assert fmt.office is kind, f".{ext}'s row must carry its module's KIND"
        assert kind.project.__module__ == kind.apply.__module__ == f"services.office.{ext}"
        assert fmt.extractor == kind.extract, "the upload reads what the edit resolves"
    assert format_for("pdf").office is None and format_for("hwpx").office is None


def test_the_superseded_code_is_gone():
    docs = (ROOT / "services" / "documents.py").read_text()
    for name in ("def extract_text_from_docx", "def extract_text_from_xlsx",
                 "def extract_text_from_pptx", "def _docx_table_text"):
        assert name not in docs, f"{name} was superseded by services/office/"
    assert not (ROOT / "services" / "export" / "office.py").exists()
    create = (ROOT / "services" / "office" / "create.py").read_text()
    rows = create[create.index("def _sheet_rows("):create.index("def _table_to_xlsx(")]
    assert '"\\\\t"' not in rows and "\\\\t" not in rows, "the literal-tab rescue existed only for the rebuild"
    assert "def write_office_file" not in create


# ── the client reads the addressed projection (executed across the line) ───


def test_client_slide_view_reads_the_addressed_projection():
    from test_adr395_model_consumable_projection import _SLIDES_PROBE, _sucrase

    if not shutil.which("node"):
        pytest.fail("node is not installed — the probe cannot run (a gate that cannot run reports nothing)")
    text, _ = opptx.project(_pptx_bytes())
    proc = subprocess.run(
        ["node", "-e", _SLIDES_PROBE, _sucrase(),
         str(ROOT.parent / "web" / "components" / "workspace" / "viewers" / "office.ts")],
        input=text, capture_output=True, text=True, timeout=60,
    )
    assert proc.returncode == 0, proc.stderr[-800:]
    slides = json.loads(proc.stdout)
    assert slides[0]["blocks"][0] == "Q3 Plan", "the address label must not reach the member"
    assert slides[0]["notes"] == "Emphasize hiring"
    assert "Metric\tValue\nARR\t$2.1M" in slides[1]["blocks"], slides[1]["blocks"]
    assert not any("[s" in b for s in slides for b in s["blocks"])
